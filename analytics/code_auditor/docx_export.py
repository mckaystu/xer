"""Export Code Analysis as a condensed Word report + priority function checklist."""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Inches, Pt, RGBColor

from analytics.code_auditor.context import contributes_to_handle_exhaustion, is_lotusscript_language
from analytics.code_auditor.models import AuditReport, Finding

# Functions that need developer attention (not SAFE / PROTECTED)
_ACTIONABLE_STATUSES = frozenset(
    {
        "UNPROTECTED_ALLOCATION",
        "PARTIAL_CLEANUP",
        "CONDITIONAL_CLEANUP",
        "ESCAPE_PATH_GAP",
    }
)

_STATUS_PRIORITY = {
    "UNPROTECTED_ALLOCATION": 0,
    "ESCAPE_PATH_GAP": 1,
    "PARTIAL_CLEANUP": 2,
    "CONDITIONAL_CLEANUP": 3,
}

_SEV_PRIORITY = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3}


def _add_checkbox_paragraph(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    run = p.add_run("☐  ")
    run.font.size = Pt(12)
    if bold_prefix:
        br = p.add_run(bold_prefix)
        br.bold = True
        br.font.size = Pt(11)
        if text:
            p.add_run(text).font.size = Pt(11)
    else:
        p.add_run(text).font.size = Pt(11)


def _sev_color(severity: str) -> RGBColor:
    s = (severity or "").upper()
    if s == "CRITICAL":
        return RGBColor(0xB9, 0x1C, 0x1C)
    if s == "HIGH":
        return RGBColor(0xC2, 0x41, 0x0C)
    if s == "MEDIUM":
        return RGBColor(0xA1, 0x62, 0x07)
    return RGBColor(0x3F, 0x3F, 0x46)


def _safe(text: str | None, limit: int = 4000) -> str:
    if not text:
        return ""
    cleaned = str(text).replace("\x00", "").strip()
    if len(cleaned) > limit:
        return cleaned[: limit - 3] + "..."
    return cleaned


def _status_label(status: str) -> str:
    return str(status or "").replace("_", " ").title()


def _actionable_functions(inventory: dict[str, Any] | None) -> list[dict[str, Any]]:
    from analytics.code_auditor.context import is_client_javascript, is_java_or_ssjs_language

    rows = []
    if isinstance(inventory, dict):
        rows = [r for r in (inventory.get("inventory") or []) if isinstance(r, dict)]
    actionable = []
    for r in rows:
        if r.get("status") not in _ACTIONABLE_STATUSES or r.get("is_false_positive"):
            continue
        lang = str(r.get("language") or r.get("language_label") or "")
        element = str(r.get("design_element") or "")
        event = str(r.get("event") or "")
        # Include Java, SSJS, and CSJS (CSJS stays LOW and sorts last).
        if is_java_or_ssjs_language(lang, event=event, element_name=element) or is_client_javascript(
            lang, event=event, element_name=element
        ):
            actionable.append(r)
    actionable.sort(key=_function_priority_key)
    return actionable


def _function_priority_key(row: dict[str, Any]) -> tuple:
    from analytics.code_auditor.context import inventory_language_priority

    sev = str(row.get("risk_severity") or row.get("severity") or "MEDIUM").upper()
    status = str(row.get("status") or "")
    in_loop = 0 if row.get("in_loop") else 1
    lang = str(row.get("language") or row.get("language_label") or "")
    element = str(row.get("design_element") or "")
    return (
        inventory_language_priority(lang, element_name=element),  # SSJS first, CSJS last
        _SEV_PRIORITY.get(sev, 9),
        _STATUS_PRIORITY.get(status, 9),
        in_loop,
        element,
        str(row.get("function_name") or ""),
    )


def _finding_sort_key(f: Finding) -> tuple:
    return (_SEV_PRIORITY.get(f.severity, 9), f.rule_id, f.element_name, f.line)


def _priority_findings(report: AuditReport, *, limit: int = 40) -> list[Finding]:
    """Short condensed finding list (CRITICAL/HIGH first) for the appendix."""
    out: list[Finding] = []
    for f in report.findings:
        if f.is_false_positive:
            continue
        if is_lotusscript_language(f.language):
            continue
        rid = f.rule_id or ""
        if contributes_to_handle_exhaustion(f) or rid.startswith(
            ("DOM-", "PERF-", "SEC-", "FORM-", "DOM-OWN", "DOM-BS")
        ):
            out.append(f)
    out.sort(key=_finding_sort_key)
    return out[:limit]


def _findings_for_function(report: AuditReport, row: dict[str, Any]) -> list[Finding]:
    """Match findings that likely belong to this inventory function."""
    name = (row.get("function_name") or "").strip().lower()
    element = (row.get("design_element") or "").strip().lower()
    if not name:
        return []
    matched: list[Finding] = []
    for f in report.findings:
        if f.is_false_positive or is_lotusscript_language(f.language):
            continue
        en = (f.element_name or "").strip().lower()
        if name in en or en == name or name in (f.title or "").lower():
            matched.append(f)
            continue
        if element and element in (f.element_name or "").lower():
            # Same design element + finding mentions function in evidence
            blob = f"{f.evidence or ''} {f.title or ''}".lower()
            if name in blob:
                matched.append(f)
    matched.sort(key=_finding_sort_key)
    return matched[:3]


def build_code_audit_checklist_docx(
    report: AuditReport,
    *,
    database_title: str | None = None,
    nsf_path: str | None = None,
    inventory: dict[str, Any] | None = None,
) -> bytes:
    """Word report: condensed analysis + priority-ordered function checklist."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    title = database_title or report.source or "Domino application"
    heading = doc.add_heading("Code Analysis — Priority Checklist", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    meta = doc.add_paragraph()
    meta.add_run("Application: ").bold = True
    meta.add_run(title)
    meta2 = doc.add_paragraph()
    meta2.add_run("NSF / source: ").bold = True
    meta2.add_run(nsf_path or report.source or "—")
    meta3 = doc.add_paragraph()
    meta3.add_run("Generated: ").bold = True
    meta3.add_run(datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))

    risk = report.risk_score()
    meta4 = doc.add_paragraph()
    meta4.add_run("Handle Exhaustion Risk: ").bold = True
    risk_run = meta4.add_run(risk)
    risk_run.bold = True
    risk_run.font.color.rgb = _sev_color(risk)

    doc.add_paragraph(
        "Scope: Java / SSJS / XPages C-API .recycle(). LotusScript is out of scope. "
        "Work the function checklist in order — each item shows the bad code to fix. "
        "Generic recycle patterns live in the Xer UI deep-dive, not repeated here."
    )

    inv_summary = {}
    if isinstance(inventory, dict):
        inv_summary = inventory.get("summary") or {}
        if not inv_summary and inventory.get("total_functions_scanned") is not None:
            inv_summary = inventory

    functions = _actionable_functions(inventory)
    findings_short = _priority_findings(report)
    exh = report.handle_exhaustion_severity_counts()
    counts = report.severity_counts()

    # —— Condensed analysis ——
    doc.add_heading("Condensed analysis", level=1)
    summary_rows = [
        ("Handle Exhaustion Risk", risk),
        (
            "Handle safety rate",
            f"{inv_summary.get('handle_safety_rate', '—')}%"
            if inv_summary.get("handle_safety_rate") is not None
            else "—",
        ),
        (
            "Functions scanned (Java/JS/XPages)",
            str(inv_summary.get("total_functions_scanned", "—")),
        ),
        ("Unprotected functions", str(inv_summary.get("unprotected_functions", "—"))),
        (
            "Partial / conditional / escape gaps",
            str(
                (inv_summary.get("functions_partial_cleanup") or 0)
                + (inv_summary.get("functions_conditional_cleanup") or 0)
                + (inv_summary.get("functions_escape_path_gap") or 0)
            ),
        ),
        (
            "Recycle coverage among allocators",
            f"{inv_summary.get('recycle_coverage_rate', '—')}%"
            if inv_summary.get("recycle_coverage_rate") is not None
            else "—",
        ),
        (
            "Handle Exhaustion CRITICAL / HIGH",
            f"{exh.get('CRITICAL', 0)} / {exh.get('HIGH', 0)}",
        ),
        (
            "Active findings (all families)",
            f"{len(report.active_findings())} · C:{counts.get('CRITICAL', 0)} "
            f"H:{counts.get('HIGH', 0)} M:{counts.get('MEDIUM', 0)} L:{counts.get('LOW', 0)}",
        ),
        ("Priority functions in checklist", str(len(functions))),
        ("Blocks scanned", f"{report.blocks_prefiltered}/{report.blocks_scanned}"),
    ]
    table = doc.add_table(rows=len(summary_rows), cols=2)
    table.style = "Table Grid"
    for i, (label, value) in enumerate(summary_rows):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = str(value)

    doc.add_paragraph()
    guide = doc.add_paragraph()
    guide.add_run("How to use: ").bold = True
    guide.add_run(
        "Address CRITICAL then HIGH functions first (especially anything in a loop). "
        "Use the Bad code block to locate the leak, apply .recycle() / finally cleanup, "
        "tick the box when fixed, then re-run Xer Code Analysis."
    )

    # —— Priority function checklist (main deliverable) ——
    doc.add_heading(f"Function checklist — priority order ({len(functions)})", level=1)
    if not functions:
        doc.add_paragraph(
            "No unprotected / partial / conditional Java/SSJS/XPages functions to checklist."
        )
    else:
        for idx, row in enumerate(functions, start=1):
            _add_function_block(doc, idx, row, report)

    # —— Condensed findings appendix ——
    doc.add_page_break()
    doc.add_heading(f"Related findings (top {len(findings_short)} by severity)", level=1)
    doc.add_paragraph(
        "Rule hits that back the function work above — evidence only (no duplicated fix templates)."
    )
    if not findings_short:
        doc.add_paragraph("No active Java/SSJS/XPages findings for this graph.")
    else:
        for idx, f in enumerate(findings_short, start=1):
            _add_finding_condensed(doc, idx, f)

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.add_run("Generated by Xer Code Analysis").italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_code_block(
    doc: Document,
    label: str,
    code: str,
    *,
    lines: list[dict[str, Any]] | None = None,
    highlight_line: int | None = None,
    limit: int = 1400,
) -> None:
    """Render a bad-code snippet; yellow-highlight the problem line(s)."""
    structured = _normalize_snippet_lines(lines, code, highlight_line=highlight_line)
    if not structured:
        return

    highlight_idxs = [i for i, row in enumerate(structured) if row.get("highlight")]
    focus = highlight_idxs[0] if highlight_idxs else 0
    kept = structured
    if sum(len(str(r.get("text") or "")) + 8 for r in structured) > limit:
        start = max(0, focus - 12)
        end = min(len(structured), focus + 13)
        kept = structured[start:end]
        while sum(len(str(r.get("text") or "")) + 8 for r in kept) > limit and len(kept) > 3:
            if focus - start >= end - focus - 1 and start < focus:
                start += 1
            elif end > focus + 1:
                end -= 1
            else:
                break
            kept = structured[start:end]

    p = doc.add_paragraph()
    p.add_run(label).bold = True

    for row in kept:
        line_no = row.get("line")
        text = str(row.get("text") or "")
        prefix = f"{line_no:>6}| " if line_no not in (None, "") else ""
        line_p = doc.add_paragraph()
        line_p.paragraph_format.space_before = Pt(0)
        line_p.paragraph_format.space_after = Pt(0)
        run = line_p.add_run(f"{prefix}{text}")
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        if row.get("highlight"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _normalize_snippet_lines(
    lines: list[dict[str, Any]] | None,
    code: str,
    *,
    highlight_line: int | None = None,
) -> list[dict[str, Any]]:
    """Build [{line, text, highlight}] from structured lines or raw snippet text."""
    out: list[dict[str, Any]] = []
    if lines:
        for row in lines:
            if not isinstance(row, dict):
                continue
            text = str(row.get("text") or "")
            abs_line = row.get("line")
            hit = bool(row.get("highlight"))
            if highlight_line and abs_line is not None and int(abs_line) == int(highlight_line):
                hit = True
            out.append({"line": abs_line, "text": text, "highlight": hit})
        if out and not any(r["highlight"] for r in out) and highlight_line:
            for r in out:
                if r.get("line") is not None and int(r["line"]) == int(highlight_line):
                    r["highlight"] = True
        if out:
            return out

    text = _safe(code, 8000)
    if not text:
        return []
    for raw in text.splitlines():
        # Numbered UI form: "  3372▶| code" or "  3372 | code"
        m = re.match(r"^\s*(\d+)\s*([▶>]?)\s*\|\s?(.*)$", raw)
        if m:
            abs_line = int(m.group(1))
            hit = bool(m.group(2)) or (
                highlight_line is not None and abs_line == int(highlight_line)
            )
            out.append({"line": abs_line, "text": m.group(3), "highlight": hit})
        else:
            out.append({"line": None, "text": raw, "highlight": False})

    if out and not any(r["highlight"] for r in out):
        if highlight_line is not None:
            for r in out:
                if r.get("line") is not None and int(r["line"]) == int(highlight_line):
                    r["highlight"] = True
        if not any(r["highlight"] for r in out) and len(out) <= 3:
            # Short evidence blob — treat whole block as the problem.
            for r in out:
                r["highlight"] = True
    return out


def _add_function_block(doc: Document, idx: int, row: dict[str, Any], report: AuditReport) -> None:
    sev = str(row.get("risk_severity") or row.get("severity") or "MEDIUM").upper()
    status = _status_label(str(row.get("status") or ""))
    name = row.get("function_name") or "function"
    element = row.get("design_element") or ""
    lang = row.get("language_label") or row.get("language") or ""
    loop_note = " · IN LOOP" if row.get("in_loop") else ""

    prefix = f"[{idx}] {sev} · {status}{loop_note}  "
    _add_checkbox_paragraph(doc, f"{name}", bold_prefix=prefix)

    loc = doc.add_paragraph()
    sev_run = loc.add_run(sev)
    sev_run.bold = True
    sev_run.font.color.rgb = _sev_color(sev)
    loc.add_run(
        f"  ·  {element}  ·  {lang}  ·  L{row.get('start_line') or row.get('line_number') or '—'}  "
        f"·  allocates {('yes' if row.get('allocates_handles') else 'no')}  "
        f"·  recycle calls {row.get('recycle_call_count', 0)}"
    )

    why = row.get("problem_breakdown") or row.get("handle_lifecycle_warning") or ""
    if why:
        p = doc.add_paragraph()
        p.add_run("Why: ").bold = True
        p.add_run(_safe(why, 500))

    unclean = row.get("unclean_vars") or []
    if unclean:
        p = doc.add_paragraph()
        p.add_run("Uncleaned handles: ").bold = True
        p.add_run(", ".join(str(v) for v in unclean[:12]))

    # Bad code only — skip generic To-Be templates (duplicated across every row).
    as_is = row.get("code_snippet_as_is") or ""
    hl = row.get("highlight_line") or row.get("line_number") or row.get("start_line")
    structured = row.get("code_snippet_lines") if isinstance(row.get("code_snippet_lines"), list) else None
    if structured and unclean:
        # Also yellow-mark allocation sites for uncleaned handle vars (not every mention).
        alloc_re = re.compile(
            r"(?:=|\bnew\b).*\b("
            + "|".join(re.escape(str(v)) for v in unclean[:12])
            + r")\b|"
            r"\b("
            + "|".join(re.escape(str(v)) for v in unclean[:12])
            + r")\b\s*=",
            re.I,
        )
        marked: list[dict[str, Any]] = []
        for item in structured:
            if not isinstance(item, dict):
                continue
            copy = dict(item)
            text = str(copy.get("text") or "")
            if alloc_re.search(text) and not re.search(r"\brecycle\s*\(|\bDelete\b", text, re.I):
                copy["highlight"] = True
            marked.append(copy)
        structured = marked
    _add_code_block(
        doc,
        "Bad code (yellow = problem line):",
        as_is,
        lines=structured,
        highlight_line=int(hl) if hl not in (None, "", 0, "0") else None,
        limit=1200,
    )

    related = _findings_for_function(report, row)
    if related:
        p = doc.add_paragraph()
        p.add_run("Related rules: ").bold = True
        p.add_run("; ".join(f"{f.rule_id} ({f.severity})" for f in related))

    notes = doc.add_paragraph()
    notes.add_run(
        "☐ Fixed   ☐ N/A / accepted risk   Owner: ___________   Date: ___________"
    ).font.size = Pt(10)
    doc.add_paragraph()


def _add_finding_condensed(doc: Document, idx: int, f: Finding) -> None:
    prefix = f"[{idx}] {f.severity} · {f.rule_id}  "
    _add_checkbox_paragraph(doc, f.title or f.rule_id, bold_prefix=prefix)
    loc = doc.add_paragraph()
    loc.add_run(
        f"{f.language_label or f.language}  ·  "
        f"{f.element_type}:{f.element_name} L{f.line}  ·  conf {f.confidence}%"
    ).font.size = Pt(10)
    bad = f.code_snippet_as_is or f.evidence or ""
    hl = f.highlight_line or f.line
    _add_code_block(
        doc,
        "Bad code (yellow = problem line):",
        bad,
        lines=f.code_snippet_lines if isinstance(f.code_snippet_lines, list) else None,
        highlight_line=int(hl) if hl else None,
        limit=900,
    )
    doc.add_paragraph()


def slug_filename(title: str | None) -> str:
    base = re.sub(r"[^\w\-]+", "_", (title or "xer_code_audit").strip())[:60].strip("_")
    return f"{base or 'xer_code_audit'}_priority_checklist.docx"


def build_code_analysis_rubric_docx() -> bytes:
    """Word doc: static search rules, scoring rubric, and AI inference passes."""
    from collections import defaultdict

    from analytics.code_auditor.models import RULE_CATALOG
    from analytics.code_auditor.snippets import PROBLEM_BREAKDOWNS, remediation_guide

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    heading = doc.add_heading("Xer Code Analysis — Rules & Rubric", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    meta = doc.add_paragraph()
    meta.add_run("Generated: ").bold = True
    meta.add_run(datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))
    meta2 = doc.add_paragraph()
    meta2.add_run(
        "This document is the live catalog of static search rules, scoring rubric, and AI "
        "inference passes used by Xer Code Analysis. It is generated from the product rule "
        "catalog (not a frozen copy)."
    )

    # —— Pipeline ——
    doc.add_heading("1. Analysis pipeline", level=1)
    for step in (
        "Extract Java / SSJS / XPages (and formula) units from the DXL / application graph.",
        "Run deterministic static search rules (DOM-*, PERF-*, SEC-*, FORM-*, DOM-OWN-*).",
        "Build the Function & Recycle Inventory (handle allocation vs cleanup per function).",
        "When OPENAI_API_KEY is set: run AI Pass 1–3 (FP filter, blind spots, ownership).",
        "Apply confidence gate (default XER_AI_CONFIDENCE_MIN=75) and human triage overrides.",
    ):
        doc.add_paragraph(step, style="List Number")

    doc.add_paragraph(
        "Handle Exhaustion scope is Java / SSJS / XPages only. LotusScript LS-DOM-* detectors "
        "remain in the catalog for reference but are not wired into Handle Exhaustion scoring."
    )

    # —— Rubric / scoring ——
    doc.add_heading("2. Scoring rubric", level=1)
    doc.add_heading("Severity", level=2)
    sev_rows = [
        ("CRITICAL", "Handle exhaustion in loops / hot paths; ODA misuse that can stall threads."),
        ("HIGH", "Missing recycle scaffolding, conditional cleanup, ownership gaps, NIF/perf risks."),
        ("MEDIUM", "Hygiene or expensive patterns that are less likely to exhaust the handle table."),
        ("LOW", "One-shot / non-loop hygiene (often demoted by AI VERIFIED_NON_LOOP)."),
    ]
    table = doc.add_table(rows=1 + len(sev_rows), cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Severity"
    table.rows[0].cells[1].text = "Meaning"
    for i, (sev, meaning) in enumerate(sev_rows, start=1):
        table.rows[i].cells[0].text = sev
        table.rows[i].cells[1].text = meaning

    doc.add_heading("Inventory statuses", level=2)
    for line in (
        "SAFE_NO_HANDLES — no Domino allocation signals in the function.",
        "PROTECTED — allocates and has matching .recycle() / Delete cleanup.",
        "PARTIAL_CLEANUP / CONDITIONAL_CLEANUP / ESCAPE_PATH_GAP — incomplete cleanup paths.",
        "UNPROTECTED_ALLOCATION — allocates with no explicit cleanup.",
    ):
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Key metrics", level=2)
    for line in (
        "Handle safety rate = (safe + fully protected) / Java-JS functions scanned.",
        "Recycle coverage among allocators = share of allocating functions that actually clean up.",
        "Allocation inside a collection loop → CRITICAL; one-shot helpers → LOW/MEDIUM hygiene.",
    ):
        doc.add_paragraph(line, style="List Bullet")

    # —— Static search rules ——
    doc.add_heading("3. Static search rules", level=1)
    doc.add_paragraph(
        "Deterministic detectors (regex / AST-light heuristics). Each finding cites a rule id."
    )

    by_cat: dict[str, list[tuple[str, dict[str, str]]]] = defaultdict(list)
    for rid, meta_r in RULE_CATALOG.items():
        by_cat[meta_r.get("category") or "Other"].append((rid, meta_r))

    preferred = [
        "C-API Handle Leaks & Object Recycling",
        "Handle Ownership",
        "Framework Conflicts",
        "Static Variables & Lifetime Anti-Patterns",
        "High-Memory & Expensive Data Patterns",
        "Performance & NIF Indexing",
        "Application Security",
        "Formula Quality",
        "AI Discrepancy & Blind Spots",
        "LotusScript Handle Lifecycle",
    ]
    categories = [c for c in preferred if c in by_cat] + sorted(
        c for c in by_cat if c not in preferred
    )

    for cat in categories:
        doc.add_heading(cat, level=2)
        if cat == "LotusScript Handle Lifecycle":
            doc.add_paragraph(
                "Reference only — not applied to Handle Exhaustion inventory or UI work list."
            )
        rules = sorted(by_cat[cat], key=lambda x: x[0])
        table = doc.add_table(rows=1 + len(rules), cols=4)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "ID"
        hdr[1].text = "Default sev"
        hdr[2].text = "Title"
        hdr[3].text = "What it searches for"
        for i, (rid, meta_r) in enumerate(rules, start=1):
            desc = PROBLEM_BREAKDOWNS.get(rid) or meta_r.get("title") or ""
            # Strip language suffix if present later — PROBLEM_BREAKDOWNS is plain.
            table.rows[i].cells[0].text = rid
            table.rows[i].cells[1].text = str(meta_r.get("default_severity") or "")
            table.rows[i].cells[2].text = str(meta_r.get("title") or "")
            table.rows[i].cells[3].text = _safe(desc, 420)
        doc.add_paragraph()

    doc.add_heading("Remediation philosophy", level=2)
    doc.add_paragraph(
        "Java / SSJS / XPages: assign intermediates, advance collections with a next-handle "
        "variable, release with .recycle() in finally on every path."
    )
    doc.add_paragraph(
        "LotusScript (hygiene reference): Delete Notes* objects (or Call obj.Recycle) before "
        "re-assignment; Set x = Nothing alone is not enough."
    )
    doc.add_paragraph(
        "ODA (org.openntf.domino): do not manually recycle — framework owns lifecycle."
    )

    # Sample guides for top handle rules
    doc.add_heading("Example remediation guides (Java)", level=2)
    for rid in ("DOM-001", "DOM-002", "DOM-010", "DOM-013", "PERF-001", "DOM-OWN-001"):
        if rid not in RULE_CATALOG:
            continue
        p = doc.add_paragraph()
        p.add_run(f"{rid}: ").bold = True
        p.add_run(_safe(remediation_guide(rid, "java"), 500))

    # —— AI inference ——
    doc.add_heading("4. AI inference rules", level=1)
    doc.add_paragraph(
        "AI runs automatically when analysis is computed and OPENAI_API_KEY is configured. "
        "Cached reloads stay fast; Refresh analysis recomputes with AI. Reviews below the "
        "confidence threshold (default 75%) are discarded."
    )

    doc.add_heading("Pass 1 — False-positive / severity filter", level=2)
    doc.add_paragraph(
        "Input: static-rule findings + surrounding code. Verdicts:"
    )
    for line in (
        "FALSE_POSITIVE — cleanup is clearly present or framework-owned (ODA, helper Delete, etc.).",
        "VERIFIED_NON_LOOP — real issue but one-shot / non-loop → demote to LOW hygiene.",
        "VERIFIED — real leak / anti-pattern, especially inside collection loops.",
    ):
        doc.add_paragraph(line, style="List Bullet")
    doc.add_paragraph(
        "Conservative: only mark FALSE_POSITIVE when cleanup is clear. Emits confidence 0–100."
    )

    doc.add_heading("Pass 2 — Blind-spot detector (DOM-BS-001)", level=2)
    doc.add_paragraph(
        "Runs on units that allocate Domino handles but had zero static hits. Looks for "
        "nested branches, early exits, conditional loops skipping recycle, exception paths, "
        "and re-assignment without releasing the prior handle. Empty result if genuinely safe."
    )

    doc.add_heading("Pass 3 — Cross-module ownership (DOM-BS-002)", level=2)
    doc.add_paragraph(
        "Caller/callee contracts across related units: who must Delete/.recycle() when a "
        "Document is returned or passed. Escalates severity for scheduled/background agents "
        "vs one-shot UI events. Skips elements already covered by deterministic DOM-OWN-001."
    )

    doc.add_heading("What AI does not do", level=2)
    for line in (
        "Does not rewrite application business logic or invent unique To-Be patches per function.",
        "Does not replace static rules — it validates, demotes, or adds residual gaps.",
        "Does not override a human Mark as false positive triage (persisted in Neon).",
    ):
        doc.add_paragraph(line, style="List Bullet")

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.add_run(
        "Generated by Xer · Download again anytime from Code Analysis to get the current catalog."
    ).italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
